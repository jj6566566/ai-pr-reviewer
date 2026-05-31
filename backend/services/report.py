from io import BytesIO
from typing import Optional
import json

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from fpdf import FPDF

from backend.models.pr_analysis import PRAnalysis

SEVERITY_LABEL = {
    "critical": "严重",
    "high": "高",
    "medium": "中",
    "low": "低",
}

RISK_LEVEL_LABEL = {
    "critical": "严重风险",
    "high": "高风险",
    "medium": "中风险",
    "low": "低风险",
}


def generate_report_md(analysis: PRAnalysis) -> str:
    lines: list[str] = []

    lines.append("# AI 代码评审报告")
    lines.append("")

    lines.append("## PR 信息")
    lines.append("")
    lines.append(f"- **标题**: {analysis.pr_title or '-'}")
    lines.append(f"- **仓库**: {analysis.repo_owner}/{analysis.repo_name}")
    lines.append(f"- **编号**: #{analysis.pr_number}")
    lines.append(f"- **作者**: {analysis.author or '-'}")
    lines.append(f"- **分支**: {analysis.head_branch or '-'} → {analysis.base_branch or '-'}")
    lines.append(f"- **文件变更**: {analysis.files_changed} 个文件 (+{analysis.additions} -{analysis.deletions})")
    lines.append(f"- **评审时间**: {analysis.created_at.strftime('%Y-%m-%d %H:%M') if analysis.created_at else '-'}")
    lines.append("")

    lines.append("## 风险评分")
    lines.append("")
    level = analysis.risk_level or "low"
    label = RISK_LEVEL_LABEL.get(level, level)
    lines.append(f"**分数**: {analysis.risk_score}/100  ({label})")
    lines.append(f"**预估评审时间**: {analysis.estimated_review_minutes} 分钟")
    lines.append("")

    if analysis.summary:
        lines.append("## 分析摘要")
        lines.append("")
        lines.append(analysis.summary.strip())
        lines.append("")

    risk_items = _parse_json(analysis.risk_items)
    if risk_items:
        lines.append("## 风险项")
        lines.append("")
        by_sev = _group_by_severity(risk_items)
        for sev in ["critical", "high", "medium", "low"]:
            items = by_sev.get(sev, [])
            if not items:
                continue
            label = SEVERITY_LABEL.get(sev, sev)
            lines.append(f"### {label} ({len(items)} 项)")
            lines.append("")
            for idx, item in enumerate(items, 1):
                file_line = f"{item.get('file', '-')}:{item.get('line', '-')}" if item.get("file") else "-"
                conf = item.get("confidence")
                conf_str = f" (置信度: {conf:.0%})" if conf is not None else ""
                lines.append(f"**{idx}. {item.get('description', '-')}**{conf_str}")
                lines.append(f"")
                lines.append(f"  📍 `{file_line}`")
                if item.get("suggestion"):
                    lines.append(f"  💡 {item.get('suggestion')}")
                lines.append("")
            lines.append("")

    suggestions = _parse_json(analysis.suggestions)
    if suggestions:
        lines.append("## 改进建议")
        lines.append("")
        for idx, sug in enumerate(suggestions, 1):
            cat = sug.get("category", "其他")
            lines.append(f"**{idx}. [{cat}] {sug.get('description', '-')}**")
            if sug.get("file"):
                lines.append(f"  📍 `{sug.get('file')}`")
            code = sug.get("code_snippet", "").strip()
            if code:
                lines.append(f"  ```")
                for cline in code.split("\n")[:10]:
                    lines.append(f"  {cline}")
                lines.append(f"  ```")
            lines.append("")

    intent = _parse_json(analysis.intent_check) if analysis.intent_check else None
    if intent and isinstance(intent, dict):
        lines.append("## 意图一致性检查")
        lines.append("")
        lines.append(f"- **一致性评分**: {intent.get('consistency_score', '-')}/100")
        lines.append(f"- **判定**: {intent.get('verdict', '-')}")
        decl = intent.get("declared_intent", "")
        actual = intent.get("actual_scope", "")
        if decl:
            lines.append(f"- **声明意图**: {decl}")
        if actual:
            lines.append(f"- **实际范围**: {actual}")
        discrepancies = intent.get("discrepancies", [])
        if discrepancies:
            lines.append("")
            lines.append("### 不一致项")
            for d in discrepancies:
                lines.append(f"- [{d.get('severity', '-')}] {d.get('description', '-')}")
        lines.append("")

    lines.append("---")
    lines.append(f"*报告由 ReviewAI 自动生成于 {analysis.created_at.strftime('%Y-%m-%d %H:%M') if analysis.created_at else '-'}*")

    return "\n".join(lines)


def _parse_json(raw: Optional[str]):
    if not raw:
        return []
    try:
        return json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return []


def _group_by_severity(items: list) -> dict:
    groups: dict[str, list] = {}
    for item in items:
        sev = item.get("severity", "low")
        groups.setdefault(sev, []).append(item)
    return groups


SEVERITY_COLORS = {
    "critical": RGBColor(0xDC, 0x26, 0x26),
    "high": RGBColor(0xF5, 0x9E, 0x0B),
    "medium": RGBColor(0x3B, 0x82, 0xF6),
    "low": RGBColor(0x6B, 0x72, 0x80),
}


def generate_report_docx(analysis: PRAnalysis) -> BytesIO:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.size = Pt(10.5)
    style.font.name = "Microsoft YaHei"

    title = doc.add_heading("AI 代码评审报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("PR 信息", level=1)
    table = doc.add_table(rows=7, cols=2, style="Light Shading Accent 1")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_rows = [
        ("标题", analysis.pr_title or "-"),
        ("仓库", f"{analysis.repo_owner}/{analysis.repo_name}"),
        ("编号", f"#{analysis.pr_number}"),
        ("作者", analysis.author or "-"),
        ("分支", f"{analysis.head_branch or '-'} → {analysis.base_branch or '-'}"),
        ("文件变更", f"{analysis.files_changed} 个文件 (+{analysis.additions} -{analysis.deletions})"),
        ("评审时间", analysis.created_at.strftime("%Y-%m-%d %H:%M") if analysis.created_at else "-"),
    ]
    for i, (key, value) in enumerate(info_rows):
        row = table.rows[i]
        row.cells[0].text = key
        row.cells[1].text = value
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    for cell in table.columns[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
        cell.width = Inches(1.5)

    doc.add_heading("风险评分", level=1)
    level = analysis.risk_level or "low"
    label = RISK_LEVEL_LABEL.get(level, level)
    p = doc.add_paragraph()
    run = p.add_run(f"分数: {analysis.risk_score}/100  ({label})")
    run.bold = True
    run.font.size = Pt(12)
    p = doc.add_paragraph(f"预估评审时间: {analysis.estimated_review_minutes} 分钟")

    if analysis.summary:
        doc.add_heading("分析摘要", level=1)
        doc.add_paragraph(analysis.summary.strip())

    risk_items = _parse_json(analysis.risk_items)
    if risk_items:
        doc.add_heading("风险项", level=1)
        by_sev = _group_by_severity(risk_items)
        for sev in ["critical", "high", "medium", "low"]:
            items = by_sev.get(sev, [])
            if not items:
                continue
            sev_label = SEVERITY_LABEL.get(sev, sev)
            color = SEVERITY_COLORS.get(sev, RGBColor(0, 0, 0))
            h = doc.add_heading(f"{sev_label} ({len(items)} 项)", level=2)
            for run in h.runs:
                run.font.color.rgb = color

            for idx, item in enumerate(items, 1):
                desc = item.get("description", "-")
                conf = item.get("confidence")
                p = doc.add_paragraph()
                run = p.add_run(f"{idx}. {desc}")
                run.bold = True
                run.font.size = Pt(10.5)
                if conf is not None:
                    run2 = p.add_run(f"  (置信度: {conf:.0%})")
                    run2.font.size = Pt(9)
                    run2.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

                file_line = f"{item.get('file', '-')}:{item.get('line', '-')}" if item.get("file") else "-"
                p2 = doc.add_paragraph(f"📍 {file_line}")
                p2.paragraph_format.left_indent = Inches(0.3)
                for r in p2.runs:
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

                suggestion = item.get("suggestion")
                if suggestion:
                    p3 = doc.add_paragraph(f"💡 {suggestion}")
                    p3.paragraph_format.left_indent = Inches(0.3)
                    for r in p3.runs:
                        r.font.size = Pt(9)

    suggestions = _parse_json(analysis.suggestions)
    if suggestions:
        doc.add_heading("改进建议", level=1)
        for idx, sug in enumerate(suggestions, 1):
            cat = sug.get("category", "其他")
            p = doc.add_paragraph()
            run = p.add_run(f"{idx}. [{cat}] {sug.get('description', '-')}")
            run.bold = True
            run.font.size = Pt(10.5)
            if sug.get("file"):
                p2 = doc.add_paragraph(f"📍 {sug.get('file')}")
                p2.paragraph_format.left_indent = Inches(0.3)
                for r in p2.runs:
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            code = sug.get("code_snippet", "").strip()
            if code:
                p_code = doc.add_paragraph()
                run_code = p_code.add_run(code[:2000])
                run_code.font.name = "Consolas"
                run_code.font.size = Pt(8)
                p_code.paragraph_format.left_indent = Inches(0.3)

    intent = _parse_json(analysis.intent_check) if analysis.intent_check else None
    if intent and isinstance(intent, dict):
        doc.add_heading("意图一致性检查", level=1)
        doc.add_paragraph(f"一致性评分: {intent.get('consistency_score', '-')}/100")
        doc.add_paragraph(f"判定: {intent.get('verdict', '-')}")
        if intent.get("declared_intent"):
            doc.add_paragraph(f"声明意图: {intent.get('declared_intent')}")
        if intent.get("actual_scope"):
            doc.add_paragraph(f"实际范围: {intent.get('actual_scope')}")
        discrepancies = intent.get("discrepancies", [])
        if discrepancies:
            doc.add_heading("不一致项", level=3)
            for d in discrepancies:
                doc.add_paragraph(f"[{d.get('severity', '-')}] {d.get('description', '-')}")

    doc.add_paragraph("")
    p_footer = doc.add_paragraph()
    run_footer = p_footer.add_run(
        f"报告由 ReviewAI 自动生成于 {analysis.created_at.strftime('%Y-%m-%d %H:%M') if analysis.created_at else '-'}"
    )
    run_footer.font.size = Pt(8)
    run_footer.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


FONT_PATH = "C:/Windows/Fonts/msyh.ttc"


def generate_report_pdf(analysis: PRAnalysis) -> BytesIO:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.add_font("yahei", "", FONT_PATH)
    pdf.add_font("yahei", "B", FONT_PATH)
    pdf.set_font("yahei", "", 10)

    pdf.set_fill_color(0xE8, 0xEC, 0xF4)
    pdf.ln(4)
    pdf.set_font("yahei", "B", 18)
    pdf.cell(0, 12, "AI 代码评审报告", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)

    pdf.set_font("yahei", "B", 13)
    pdf.set_text_color(0x1E, 0x24, 0x40)
    pdf.cell(0, 8, "PR 信息", new_x="LMARGIN", new_y="NEXT")
    pdf.set_draw_color(0x7B, 0x82, 0x9C)
    pdf.line(pdf.l_margin, pdf.get_y() + 1, pdf.w - pdf.r_margin, pdf.get_y() + 1)
    pdf.ln(4)

    info_pairs = [
        ("标题", analysis.pr_title or "-"),
        ("仓库", f"{analysis.repo_owner}/{analysis.repo_name}"),
        ("编号", f"#{analysis.pr_number}"),
        ("作者", analysis.author or "-"),
        ("分支", f"{analysis.head_branch or '-'}  →  {analysis.base_branch or '-'}"),
        ("文件变更", f"{analysis.files_changed} 个文件 (+{analysis.additions} -{analysis.deletions})"),
        ("评审时间", analysis.created_at.strftime("%Y-%m-%d %H:%M") if analysis.created_at else "-"),
    ]
    for key, val in info_pairs:
        pdf.set_text_color(0x58, 0x5E, 0x74)
        pdf.set_font("yahei", "", 9)
        pdf.cell(22, 6, key + ":")
        pdf.set_text_color(0x1A, 0x1F, 0x36)
        pdf.set_font("yahei", "", 9)
        pdf.cell(0, 6, val, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    pdf.set_font("yahei", "B", 13)
    pdf.set_text_color(0x1E, 0x24, 0x40)
    pdf.cell(0, 8, "风险评分", new_x="LMARGIN", new_y="NEXT")
    pdf.line(pdf.l_margin, pdf.get_y() + 1, pdf.w - pdf.r_margin, pdf.get_y() + 1)
    pdf.ln(4)

    level = analysis.risk_level or "low"
    label = RISK_LEVEL_LABEL.get(level, level)
    pdf.set_font("yahei", "B", 12)
    pdf.set_text_color(0x1A, 0x1F, 0x36)
    pdf.cell(0, 7, f"分数: {analysis.risk_score}/100  ({label})", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("yahei", "", 9)
    pdf.set_text_color(0x58, 0x5E, 0x74)
    pdf.cell(0, 6, f"预估评审时间: {analysis.estimated_review_minutes} 分钟", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    if analysis.summary:
        pdf.set_font("yahei", "B", 13)
        pdf.set_text_color(0x1E, 0x24, 0x40)
        pdf.cell(0, 8, "分析摘要", new_x="LMARGIN", new_y="NEXT")
        pdf.line(pdf.l_margin, pdf.get_y() + 1, pdf.w - pdf.r_margin, pdf.get_y() + 1)
        pdf.ln(4)
        pdf.set_font("yahei", "", 9)
        pdf.set_text_color(0x1A, 0x1F, 0x36)
        pdf.multi_cell(0, 5, analysis.summary.strip())
        pdf.ln(3)

    risk_items = _parse_json(analysis.risk_items)
    if risk_items:
        pdf.set_font("yahei", "B", 13)
        pdf.set_text_color(0x1E, 0x24, 0x40)
        pdf.cell(0, 8, "风险项", new_x="LMARGIN", new_y="NEXT")
        pdf.line(pdf.l_margin, pdf.get_y() + 1, pdf.w - pdf.r_margin, pdf.get_y() + 1)
        pdf.ln(5)

        by_sev = _group_by_severity(risk_items)
        for sev in ["critical", "high", "medium", "low"]:
            items = by_sev.get(sev, [])
            if not items:
                continue
            sev_label = SEVERITY_LABEL.get(sev, sev)
            r, g, b = SEVERITY_COLORS[sev]
            pdf.set_font("yahei", "B", 11)
            pdf.set_text_color(r / 255 * 255, g / 255 * 255, b / 255 * 255)
            pdf.cell(0, 7, f"{sev_label} ({len(items)} 项)", new_x="LMARGIN", new_y="NEXT")
            pdf.set_text_color(0x1A, 0x1F, 0x36)

            for idx, item in enumerate(items, 1):
                desc = item.get("description", "-")
                conf = item.get("confidence")
                y_before = pdf.get_y()
                if y_before > 260:
                    pdf.add_page()

                pdf.set_font("yahei", "B", 9)
                header = f"  {idx}. {desc}"
                if conf is not None:
                    header += f"  (置信度: {conf:.0%})"
                pdf.multi_cell(0, 5, header)

                file_line = f"{item.get('file', '-')}:{item.get('line', '-')}" if item.get("file") else "-"
                pdf.set_font("yahei", "", 8)
                pdf.set_text_color(0x58, 0x5E, 0x74)
                pdf.cell(8, 5, "")
                pdf.cell(0, 5, f"📍 {file_line}", new_x="LMARGIN", new_y="NEXT")

                suggestion = item.get("suggestion")
                if suggestion:
                    pdf.cell(8, 5, "")
                    pdf.multi_cell(0, 5, f"💡 {suggestion}")
                pdf.set_text_color(0x1A, 0x1F, 0x36)
                pdf.ln(1)
            pdf.ln(2)

    suggestions = _parse_json(analysis.suggestions)
    if suggestions:
        pdf.add_page()
        pdf.set_font("yahei", "B", 13)
        pdf.set_text_color(0x1E, 0x24, 0x40)
        pdf.cell(0, 8, "改进建议", new_x="LMARGIN", new_y="NEXT")
        pdf.line(pdf.l_margin, pdf.get_y() + 1, pdf.w - pdf.r_margin, pdf.get_y() + 1)
        pdf.ln(5)

        for idx, sug in enumerate(suggestions, 1):
            cat = sug.get("category", "其他")
            pdf.set_font("yahei", "B", 9)
            pdf.set_text_color(0x1A, 0x1F, 0x36)
            pdf.cell(0, 5, f"  {idx}. [{cat}] {sug.get('description', '-')}", new_x="LMARGIN", new_y="NEXT")
            if sug.get("file"):
                pdf.set_font("yahei", "", 8)
                pdf.set_text_color(0x58, 0x5E, 0x74)
                pdf.cell(8, 5, "")
                pdf.cell(0, 5, f"📍 {sug.get('file')}", new_x="LMARGIN", new_y="NEXT")
            code = sug.get("code_snippet", "").strip()
            if code:
                pdf.set_font("yahei", "", 7)
                pdf.set_text_color(0x1A, 0x1F, 0x36)
                for cline in code.split("\n")[:8]:
                    if cline.strip():
                        pdf.cell(10, 4, "")
                        pdf.cell(0, 4, f"  {cline[:120]}", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1.5)
        pdf.ln(2)

    intent = _parse_json(analysis.intent_check) if analysis.intent_check else None
    if intent and isinstance(intent, dict):
        pdf.set_font("yahei", "B", 13)
        pdf.set_text_color(0x1E, 0x24, 0x40)
        pdf.cell(0, 8, "意图一致性检查", new_x="LMARGIN", new_y="NEXT")
        pdf.line(pdf.l_margin, pdf.get_y() + 1, pdf.w - pdf.r_margin, pdf.get_y() + 1)
        pdf.ln(4)
        pdf.set_font("yahei", "", 9)
        pdf.set_text_color(0x1A, 0x1F, 0x36)
        pdf.cell(0, 6, f"一致性评分: {intent.get('consistency_score', '-')}/100", new_x="LMARGIN", new_y="NEXT")
        pdf.cell(0, 6, f"判定: {intent.get('verdict', '-')}", new_x="LMARGIN", new_y="NEXT")
        if intent.get("declared_intent"):
            pdf.cell(0, 6, f"声明意图: {intent.get('declared_intent')}", new_x="LMARGIN", new_y="NEXT")
        if intent.get("actual_scope"):
            pdf.cell(0, 6, f"实际范围: {intent.get('actual_scope')}", new_x="LMARGIN", new_y="NEXT")
        discrepancies = intent.get("discrepancies", [])
        if discrepancies:
            pdf.ln(2)
            pdf.set_font("yahei", "B", 10)
            pdf.cell(0, 6, "  不一致项:", new_x="LMARGIN", new_y="NEXT")
            for d in discrepancies:
                pdf.set_font("yahei", "", 9)
                pdf.cell(0, 6, f"    [{d.get('severity', '-')}] {d.get('description', '-')}", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(4)

    pdf.set_font("yahei", "", 7)
    pdf.set_text_color(0xA0, 0xA6, 0xB8)
    ts = analysis.created_at.strftime("%Y-%m-%d %H:%M") if analysis.created_at else "-"
    pdf.cell(0, 10, f"报告由 ReviewAI 自动生成于 {ts}", align="C")

    buf = BytesIO()
    buf.write(pdf.output())
    buf.seek(0)
    return buf
